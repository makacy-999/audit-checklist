# -*- coding: utf-8 -*-
"""生成 index.html 与 Excel 版 Check list"""
import json, os, sys
from collections import OrderedDict, defaultdict

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.path.insert(0, HERE)
from data import DATA, PERSON_DEPT, DEPT_OPTIONS

# ---------- 1. index.html ----------
tpl = open(os.path.join(HERE, 'template.html'), encoding='utf-8').read()
html = (tpl
        .replace('__DATA__', json.dumps(DATA, ensure_ascii=False))
        .replace('__DEPTS__', json.dumps(DEPT_OPTIONS, ensure_ascii=False))
        .replace('__PDEPT__', json.dumps(PERSON_DEPT, ensure_ascii=False)))
with open(os.path.join(ROOT, 'index.html'), 'w', encoding='utf-8') as f:
    f.write(html)
print('index.html ok', len(html))

# ---------- 2. flatten ----------
FLAT = []
for si, sec in enumerate(DATA):
    for it in sec['items']:
        bo, bd = it.get('owner', ''), it.get('due', '')
        if not it['files']:
            FLAT.append(dict(cat=sec['cat'], id=it['id'], zh=it['zh'], en=it['en'],
                             cat_en=it.get('cat_en', ''), g='', text='(本条无文件清单)',
                             owner=bo, due=bd, nofile=True, subs=it.get('subs', [])))
        for f in it['files']:
            FLAT.append(dict(cat=sec['cat'], id=it['id'], zh=it['zh'], en=it['en'],
                             cat_en=it.get('cat_en', ''), g=f.get('g', ''), text=f['t'],
                             owner=f.get('owner') or bo, due=f.get('due') or bd,
                             nofile=False, subs=it.get('subs', [])))

def dept_of(owner):
    if not owner:
        return ''
    first = owner.split('、')[0].strip()
    return PERSON_DEPT.get(first) or PERSON_DEPT.get(owner) or ''

# ---------- 3. Excel ----------
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

THIN = Side(style='thin', color='D5DBE4')
BD = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HDR = PatternFill('solid', fgColor='1A56A7')
HF = Font(color='FFFFFF', bold=True, size=10.5, name='微软雅黑')
BF = Font(size=10.5, name='微软雅黑')
BF2 = Font(size=10, name='微软雅黑', color='4A586B')
WRAP = Alignment(wrap_text=True, vertical='top')
CTR = Alignment(horizontal='center', vertical='center', wrap_text=True)

wb = openpyxl.Workbook()

# --- Sheet 1 逐项核对 ---
ws = wb.active
ws.title = '① 逐项核对清单'
heads = ['分类', '编号', '审计主题（中文）', '英文原文', '分组', '准备文件清单（逐项核对）',
         '责任部门', '责任人', '完成时间', '状态', '备注']
widths = [26, 7, 34, 44, 8, 46, 15, 20, 12, 11, 24]
ws.append(heads)
for i, w in enumerate(widths, 1):
    ws.column_dimensions[get_column_letter(i)].width = w
for c in ws[1]:
    c.fill, c.font, c.alignment, c.border = HDR, HF, CTR, BD
ws.row_dimensions[1].height = 30

CATFILL = PatternFill('solid', fgColor='EAF1FB')
r = 2
prev_cat = None
for f in FLAT:
    ws.append([f['cat'], f['id'], f['zh'], (('[' + f['cat_en'] + '] ') if f['cat_en'] else '') + f['en'],
               f['g'], f['text'], dept_of(f['owner']), f['owner'] or '—',
               f['due'] or '—', '未开始', ''])
    for c in ws[r]:
        c.font, c.alignment, c.border = BF, WRAP, BD
    if f['cat'] != prev_cat:
        for c in ws[r][:3]:
            c.fill = CATFILL
        ws[r][0].font = Font(size=10.5, bold=True, name='微软雅黑', color='1A56A7')
        prev_cat = f['cat']
    ws[r][1].alignment = CTR
    ws[r][4].alignment = CTR
    ws[r][8].alignment = CTR
    ws[r][9].alignment = CTR
    if f['nofile']:
        ws[r][5].font = Font(size=10, name='微软雅黑', color='8492A6', italic=True)
    r += 1

dv = DataValidation(type='list', formula1='"未开始,准备中,已就绪,不适用"', allow_blank=True)
dv.error = '请从下拉列表选择'
dv.prompt = '未开始 / 准备中 / 已就绪 / 不适用'
ws.add_data_validation(dv)
dv.add(f'J2:J{r-1}')

dv2 = DataValidation(type='list',
                     formula1='"' + ','.join(DEPT_OPTIONS) + '"', allow_blank=True)
ws.add_data_validation(dv2)
dv2.add(f'G2:G{r-1}')

ws.freeze_panes = 'C2'
ws.auto_filter.ref = f'A1:K{r-1}'

# --- Sheet 2 按责任人派工 ---
ws2 = wb.create_sheet('② 按责任人派工')
by = defaultdict(list)
for f in FLAT:
    if not f['owner'] or f['owner'] == '/':
        by['（未指定）'].append(f)
        continue
    for p in f['owner'].split('、'):
        by[p.strip()].append(f)

ws2.append(['责任人', '责任部门', '核对项数', '编号', '准备文件清单', '审计主题（中文）', '完成时间', '状态', '备注'])
for i, w in enumerate([12, 20, 10, 7, 50, 34, 12, 11, 24], 1):
    ws2.column_dimensions[get_column_letter(i)].width = w
for c in ws2[1]:
    c.fill, c.font, c.alignment, c.border = HDR, HF, CTR, BD
ws2.row_dimensions[1].height = 28

names = [n for n in by if n != '（未指定）']
names.sort(key=lambda x: (dept_of(x), x))
if '（未指定）' in by:
    names.append('（未指定）')
r2 = 2
for n in names:
    rows = by[n]
    start = r2
    for f in rows:
        ws2.append([n, dept_of(n) or '—', '', f['id'], f['text'], f['zh'],
                    f['due'] or '—', '未开始', ''])
        for c in ws2[r2]:
            c.font, c.alignment, c.border = BF, WRAP, BD
        ws2[r2][3].alignment = CTR
        ws2[r2][6].alignment = CTR
        ws2[r2][7].alignment = CTR
        r2 += 1
    ws2.cell(start, 1).font = Font(size=10.5, bold=True, name='微软雅黑')
    ws2.cell(start, 3).value = len(rows)
    ws2.cell(start, 3).alignment = CTR
    ws2.cell(start, 3).font = Font(size=10.5, bold=True, name='微软雅黑', color='1A56A7')
    if len(rows) > 1:
        ws2.merge_cells(start_row=start, start_column=1, end_row=r2-1, end_column=1)
        ws2.merge_cells(start_row=start, start_column=2, end_row=r2-1, end_column=2)
        ws2.merge_cells(start_row=start, start_column=3, end_row=r2-1, end_column=3)
        ws2.cell(start, 1).alignment = CTR
        ws2.cell(start, 2).alignment = CTR

dv3 = DataValidation(type='list', formula1='"未开始,准备中,已就绪,不适用"', allow_blank=True)
ws2.add_data_validation(dv3)
dv3.add(f'H2:H{r2-1}')
dv4 = DataValidation(type='list', formula1='"' + ','.join(DEPT_OPTIONS) + '"', allow_blank=True)
ws2.add_data_validation(dv4)
dv4.add(f'B2:B{r2-1}')
ws2.freeze_panes = 'D2'
ws2.auto_filter.ref = f'A1:I{r2-1}'

# --- Sheet 3 说明 ---
ws3 = wb.create_sheet('③ 使用说明')
ws3.column_dimensions['A'].width = 22
ws3.column_dimensions['B'].width = 96
notes = [
    ('文件用途', '客户审计应对逐项核对清单。按审计日程原文分类，将每项「准备资料」拆成可单独勾选的子项，并补齐责任部门与责任人，供质量管理人员逐条核对、防止遗漏。'),
    ('数据来源', '《09:17审计日程和需提供的文件清单.xlsx》——「第一天审计日程」与「附件-需提供文件清单」两个工作表。'),
    ('处理原则', '原文描述与顺序未作改写；仅做三件事：①按审计模块分类；②把「准备资料」中的 ①②③ 拆成独立核对项；③补责任部门。'),
    ('责任部门', '原文只给了人名。部门是按每人承担的工作推断的预填值，可直接在「责任部门」列下拉修改。'),
    ('怎么用', 'Sheet①：质量管理人员主控，逐条改状态、填备注，可按分类/责任人筛选。\nSheet②：按人拆开，直接截图或打印发给对应同事。\n网页版 index.html：进度自动统计、可搜索筛选、打印存 PDF、导出 CSV。'),
    ('状态口径', '未开始＝还没动手；准备中＝整理中；已就绪＝可随时调阅/提交；不适用＝本次审计不涉及。'),
    ('重点风险', '1) 2.5 温控车温度分布验证报告须先经研发部审核后，方可提交客户查看，未经审核不得外发。\n2) 2.6/5.1/5.3 涉及系统说明与现场演示，需提前演练。\n3) 第八类「附件-需提供文件清单」中多项（近两年批次清单、OOS清单、召回清单、退货清单）原文未指定责任人，需尽快指派。'),
    ('截止时间', '多数条目截止 9月11日前；3.7 审计与自检 9月12日前；3.8 分包商管理 9月13日前；2.6 第2组（异常处理记录、监控岗日志）9月14日。'),
]
ws3.append(['项目', '说明'])
for c in ws3[1]:
    c.fill, c.font, c.alignment, c.border = HDR, HF, CTR, BD
for i, (k, v) in enumerate(notes, start=2):
    ws3.append([k, v])
    ws3.cell(i, 1).font = Font(size=10.5, bold=True, name='微软雅黑', color='1A56A7')
    ws3.cell(i, 1).alignment = Alignment(vertical='top', wrap_text=True)
    ws3.cell(i, 2).font = BF
    ws3.cell(i, 2).alignment = WRAP
    ws3.cell(i, 1).border = BD
    ws3.cell(i, 2).border = BD
    ws3.row_dimensions[i].height = max(30, 15 * (v.count('\n') + 1 + len(v) // 60))

out = os.path.join(ROOT, '审计应对Checklist.xlsx')
wb.save(out)

print('rows FLAT =', len(FLAT))
print('persons =', len(by))
print('xlsx ->', out)
